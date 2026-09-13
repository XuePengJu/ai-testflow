"""文档抽取链路测试（docx / pdf / md / txt）。

覆盖三件事：
1. `doc_extract.extract_text` 能读出各格式文本，不支持格式抛明确异常；
2. 解析链路（parser_agent.run_parser）对 docx 不再崩（旧行为：read_text 抛
   UnicodeDecodeError → 任务 failed）；
3. 端到端：POST /api/files 上传 docx 拿到 file_id，对话时能读回文档文本。
"""
from io import BytesIO

import pytest

from app.services.doc_extract import (
    ExtractError,
    UnsupportedFormatError,
    extract_text,
    extract_text_safe,
    is_supported,
    supported_hint,
)


# ---------- 夹具构造 ----------

def _make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    """用 python-docx 现场生成一个 docx（含段落，可选表格）。"""
    import docx

    d = docx.Document()
    for p in paragraphs:
        d.add_paragraph(p)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for i, row in enumerate(table):
            for j, cell in enumerate(row):
                t.cell(i, j).text = cell
    buf = BytesIO()
    d.save(buf)
    return buf.getvalue()


def _make_pdf(text: str) -> bytes:
    """构造最小可解析 PDF（单页、Helvetica、纯 ASCII）。"""
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
    ]
    content = f"BT /F1 24 Tf 72 700 Td ({text}) Tj ET".encode("latin-1")
    objs.append(
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream"
    )
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, o in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + o + b"\nendobj\n"
    xref_pos = len(out)
    out += f"xref\n0 {len(objs) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF\n"
    ).encode()
    return bytes(out)


# ---------- 白名单 ----------

def test_supported_exts():
    assert is_supported("需求文档.docx")
    assert is_supported("spec.PDF")          # 大小写不敏感
    assert is_supported("notes.md")
    assert is_supported("readme.markdown")
    assert is_supported("plain.txt")
    assert not is_supported("报表.xlsx")
    assert not is_supported("data.json")
    assert not is_supported("无扩展名")


def test_supported_hint_lists_all():
    hint = supported_hint()
    for ext in ("docx", "pdf", "md", "markdown", "txt"):
        assert ext in hint


# ---------- 各格式抽取 ----------

def test_extract_txt(tmp_path):
    p = tmp_path / "req.txt"
    p.write_text("登录功能需求\n支持账号密码登录", encoding="utf-8")
    assert extract_text(p) == "登录功能需求\n支持账号密码登录"


def test_extract_md(tmp_path):
    p = tmp_path / "req.md"
    p.write_text("## 登录\n\n输入正确的账号密码可登录", encoding="utf-8")
    assert "输入正确的账号密码可登录" in extract_text(p)


def test_extract_docx_paragraphs(tmp_path):
    p = tmp_path / "req.docx"
    p.write_bytes(_make_docx(["采购管理需求", "新增采购单后可提交审批"]))
    text = extract_text(p)
    assert "采购管理需求" in text
    assert "新增采购单后可提交审批" in text


def test_extract_docx_tables(tmp_path):
    """需求文档的规则常在表格里，必须一并抽出来。"""
    p = tmp_path / "req.docx"
    p.write_bytes(_make_docx(["字段规则"], table=[["字段", "规则"], ["金额", "超 5 万需二级审批"]]))
    text = extract_text(p)
    assert "金额" in text and "超 5 万需二级审批" in text
    assert "|" in text          # 表格按 | 拼接


def test_extract_pdf(tmp_path):
    p = tmp_path / "spec.pdf"
    p.write_bytes(_make_pdf("Login Requirement Specification"))
    assert "Login Requirement" in extract_text(p)


def test_extract_unsupported_raises(tmp_path):
    p = tmp_path / "报表.xlsx"
    p.write_bytes(b"PK\x03\x04")
    with pytest.raises(UnsupportedFormatError):
        extract_text(p)


def test_extract_broken_docx_raises_extract_error(tmp_path):
    p = tmp_path / "bad.docx"
    p.write_bytes(b"not a real docx")
    with pytest.raises(ExtractError):
        extract_text(p)


def test_extract_text_safe_never_raises(tmp_path):
    """安全版：任何失败都返回空串，绝不让解析链路崩。"""
    bad = tmp_path / "报表.xlsx"
    bad.write_bytes(b"PK\x03\x04")
    assert extract_text_safe(bad) == ""
    assert extract_text_safe(tmp_path / "不存在.docx") == ""


def test_binary_was_previously_fatal(tmp_path):
    """回归锚点：旧实现直接 read_text(utf-8) 读二进制会抛 UnicodeDecodeError。"""
    p = tmp_path / "req.docx"
    p.write_bytes(_make_docx(["边界场景需求"]))
    with pytest.raises(UnicodeDecodeError):
        p.read_text(encoding="utf-8")
    assert "边界场景需求" in extract_text(p)      # 新链路正常读出


# ---------- 解析链路接线 ----------

def test_run_parser_reads_docx(tmp_path):
    """docx 输入能走到解析出测试点（无 AI 时走正则/整段兜底），不再抛异常。"""
    from app.workflow.agents.parser_agent import run_parser

    p = tmp_path / "需求.docx"
    p.write_bytes(_make_docx(["## 登录功能", "账号密码登录", "## 权限控制", "不同角色可见范围不同"]))
    units, summary, details = run_parser(str(p), "business", client=None)
    assert units, "docx 应当能解析出测试点"
    assert "测试点" in summary
    import json

    meta = json.loads(details)
    assert meta.get("title")


# ---------- 端到端：上传 + 对话注入 ----------

def test_upload_docx_end_to_end(client, accounts):
    tok = accounts["user"]["token"]
    r = client.post(
        "/api/files",
        files={"file": ("需求文档.docx", _make_docx(["电商下单主流程需求"]), "application/octet-stream")},
        headers={"Authorization": "Bearer " + tok},
    )
    assert r.status_code == 200, r.text
    data = r.json()
    assert data["file_id"]
    assert data["chars"] > 0
    assert data["name"] == "需求文档.docx"

    # 读回：对话端点应能按 file_id 取到文档文本
    from app.api.chat import _build_attachment_context
    from app.api.files import load_chat_file

    loaded = load_chat_file(data["file_id"])
    assert loaded is not None
    assert "电商下单主流程需求" in loaded[1]
    ctx = _build_attachment_context(loaded)
    assert "需求文档.docx" in ctx and "电商下单主流程需求" in ctx


def test_upload_pdf_end_to_end(client, accounts):
    tok = accounts["user"]["token"]
    r = client.post(
        "/api/files",
        files={"file": ("spec.pdf", _make_pdf("Order Flow Requirement"), "application/pdf")},
        headers={"Authorization": "Bearer " + tok},
    )
    assert r.status_code == 200, r.text
    assert r.json()["chars"] > 0


def test_upload_rejects_unsupported(client, accounts):
    tok = accounts["user"]["token"]
    r = client.post(
        "/api/files",
        files={"file": ("报表.xlsx", b"PK\x03\x04", "application/octet-stream")},
        headers={"Authorization": "Bearer " + tok},
    )
    assert r.status_code == 400
    assert "不支持" in r.text or "格式" in r.text


def test_chat_with_attachment_reads_doc(client, accounts):
    """带 file_id 的对话：mock 回复里应体现已读取附件（不再反问要需求）。"""
    tok = accounts["user"]["token"]
    up = client.post(
        "/api/files",
        files={"file": ("需求.docx", _make_docx(["库存盘点需求"]), "application/octet-stream")},
        headers={"Authorization": "Bearer " + tok},
    )
    assert up.status_code == 200, up.text
    fid = up.json()["file_id"]

    r = client.post(
        "/api/chat/stream",
        json={"message": "", "file_id": fid},
        headers={"Authorization": "Bearer " + tok},
    )
    assert r.status_code == 200, r.text
    assert "已读取附件" in r.text
    assert "需求.docx" in r.text


def test_create_task_rejects_unsupported_file(client, accounts):
    """/api/tasks 也要拦：不支持格式提前 400，而不是走到解析步骤才崩。"""
    tok = accounts["user"]["token"]
    r = client.post(
        "/api/tasks",
        files={"file": ("报表.xlsx", b"PK\x03\x04", "application/octet-stream")},
        data={"text": "", "kind": "business"},
        headers={"Authorization": "Bearer " + tok},
    )
    assert r.status_code == 400
    assert "不支持" in r.text


def test_create_task_accepts_docx(client, accounts):
    """docx 走完整任务创建（不触发后台执行断言，仅验证不再被拦）。"""
    tok = accounts["user"]["token"]
    r = client.post(
        "/api/tasks",
        files={"file": ("需求.docx", _make_docx(["用户注册需求"]), "application/octet-stream")},
        data={"text": "", "kind": "business", "formats": "xlsx"},
        headers={"Authorization": "Bearer " + tok},
    )
    assert r.status_code == 201, r.text
    assert r.json()["source_type"] == "file"
