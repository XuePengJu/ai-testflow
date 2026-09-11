"""ImportAgent：导入已有用例文件，解析为结构化 TestCase 列表。

支持格式：
- xmind：平台导出的 XMind 8 legacy（content.xml）+ 新版 XMind 2020+（content.json）
- xlsx：平台导出格式（用例ID/标题/模块/类型/优先级/前置条件/步骤/预期结果/测试数据），兼容常见列名别名
- json：平台导出的 TestCase 数组，兼容字段别名

解析失败时抛出 ImportError，不静默降级。
返回 (cases, summary)，summary 含总数、模块分布、类型分布。
"""
import json
import zipfile
from collections import Counter
from xml.etree import ElementTree as ET

from src.models.testcase import TestCase, CaseType, Priority, align_step_expectations


class ImportError(Exception):
    """用例文件解析失败。"""


# ============ 类型/优先级归一化 ============

_TYPE_MAP = {
    "正向": "正向", "positive": "正向", "正常": "正向", "正向用例": "正向",
    "异常": "异常", "negative": "异常", "出错": "异常", "异常用例": "异常",
    "边界": "边界值", "边界值": "边界值", "boundary": "边界值", "边界值用例": "边界值",
    "场景": "场景组合", "场景组合": "场景组合", "scenario": "场景组合", "场景组合用例": "场景组合",
}


def _norm_type(ct: str) -> str:
    if not ct:
        return "正向"
    ct = str(ct).strip()
    return _TYPE_MAP.get(ct, ct)


def _to_case_type(ct: str) -> CaseType:
    t = _norm_type(ct)
    if t == "正向":
        return CaseType.POSITIVE
    if t == "异常":
        return CaseType.NEGATIVE
    if t == "边界值":
        return CaseType.BOUNDARY
    return CaseType.SCENARIO


def _to_priority(pr: str) -> Priority:
    pr = str(pr or "P1").strip().upper()
    for p in ("P0", "P1", "P2", "P3"):
        if pr.startswith(p):
            return Priority(p)
    return Priority.P1


# ============ JSON 导入 ============

def import_json(path: str) -> list[TestCase]:
    """解析平台导出的 json 用例文件。"""
    try:
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
    except (OSError, json.JSONDecodeError) as e:
        raise ImportError(f"JSON 文件读取失败：{e}") from e
    if not isinstance(data, list):
        raise ImportError("JSON 顶层必须是用例数组")
    cases = []
    for i, item in enumerate(data):
        if not isinstance(item, dict):
            continue
        try:
            steps = item.get("steps") or item.get("步骤") or []
            expected = str(item.get("expected") or item.get("预期结果") or "")
            se = item.get("step_expectations") or item.get("步骤预期") or []
            cases.append(TestCase(
                case_id=str(item.get("case_id") or item.get("用例ID") or ""),
                title=str(item.get("title") or item.get("标题") or "未命名用例"),
                module=str(item.get("module") or item.get("模块") or ""),
                case_type=_to_case_type(item.get("case_type") or item.get("类型") or "正向"),
                priority=_to_priority(item.get("priority") or item.get("优先级") or "P1"),
                pre_condition=str(item.get("pre_condition") or item.get("前置条件") or ""),
                steps=steps,
                step_expectations=align_step_expectations(steps, se, expected),
                expected=expected,
                test_data=item.get("test_data") or item.get("测试数据"),
            ))
        except Exception as e:  # noqa: BLE001
            raise ImportError(f"第 {i + 1} 条用例解析失败：{e}") from e
    if not cases:
        raise ImportError("JSON 中未解析到有效用例")
    return cases


# ============ XLSX 导入 ============

# 列名别名映射：标准字段 -> 可能的列名列表
_COL_ALIASES = {
    "case_id": ["用例id", "case_id", "id", "编号", "用例编号"],
    "title": ["标题", "title", "用例标题", "名称"],
    "module": ["模块", "module", "所属模块"],
    "case_type": ["类型", "case_type", "type", "用例类型"],
    "priority": ["优先级", "priority"],
    "pre_condition": ["前置条件", "pre_condition", "precondition", "前置"],
    "steps": ["步骤", "steps", "操作步骤", "测试步骤"],
    "step_expectations": ["步骤预期", "step_expectations", "步骤预期结果"],
    "expected": ["预期结果", "expected", "预期"],
    "test_data": ["测试数据", "test_data", "testdata", "数据"],
}


def _match_col(header: str, aliases: list[str]) -> bool:
    h = str(header).strip().lower().replace(" ", "")
    return h in aliases


def import_xlsx(path: str) -> list[TestCase]:
    """解析平台导出的 xlsx 用例文件。"""
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise ImportError("缺少 openpyxl 依赖，无法解析 xlsx") from e
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
    except Exception as e:  # noqa: BLE001
        raise ImportError(f"xlsx 读取失败：{e}") from e
    if not rows:
        raise ImportError("xlsx 为空")
    # 第一行是表头
    headers = [str(c or "").strip() for c in rows[0]]
    # 建立 标准字段 -> 列索引 映射
    col_idx: dict[str, int] = {}
    for std, aliases in _COL_ALIASES.items():
        for i, h in enumerate(headers):
            if _match_col(h, aliases):
                col_idx[std] = i
                break
    if "title" not in col_idx:
        raise ImportError(f"xlsx 表头未找到「标题」列，现有表头：{headers}")
    cases = []
    for row in rows[1:]:
        if not row or all(c is None or str(c).strip() == "" for c in row):
            continue
        def _get(key: str) -> str:
            idx = col_idx.get(key)
            if idx is None or idx >= len(row):
                return ""
            v = row[idx]
            return "" if v is None else str(v).strip()
        steps_raw = _get("steps")
        steps = [s.strip() for s in steps_raw.split("\n") if s.strip()] if steps_raw else []
        se_raw = _get("step_expectations")
        se = [s.strip() for s in se_raw.split("\n") if s.strip()] if se_raw else []
        expected = _get("expected")
        try:
            cases.append(TestCase(
                case_id=_get("case_id"),
                title=_get("title") or "未命名用例",
                module=_get("module"),
                case_type=_to_case_type(_get("case_type") or "正向"),
                priority=_to_priority(_get("priority") or "P1"),
                pre_condition=_get("pre_condition"),
                steps=steps,
                step_expectations=align_step_expectations(steps, se, expected),
                expected=expected,
                test_data=_get("test_data") or None,
            ))
        except Exception as e:  # noqa: BLE001
            raise ImportError(f"xlsx 行解析失败：{e}") from e
    if not cases:
        raise ImportError("xlsx 中未解析到有效用例")
    return cases


# ============ XMind 导入 ============

def _strip_step_num(title: str) -> str:
    """去掉操作步骤标题前的 `1. ` 编号。"""
    import re
    return re.sub(r"^\d+\.\s*", "", title or "").strip()


def _parse_case_topic_xml(topic: ET.Element, ns: str) -> TestCase:
    """从 XML topic 元素还原一个 TestCase。"""
    title_el = topic.find(f"{ns}title")
    raw_title = (title_el.text or "").strip() if title_el is not None else ""
    # 标题格式：`序号. 操作概括->预期概括`，去掉序号和 -> 后缀
    import re
    clean_title = re.sub(r"^\d+\.\s*", "", raw_title)
    clean_title = re.sub(r"->.*$", "", clean_title).strip()

    # labels：[类型用例, 优先级]
    case_type = "正向"
    priority = "P1"
    labels_el = topic.find(f"{ns}labels")
    if labels_el is not None:
        for label in labels_el.findall(f"{ns}label"):
            t = (label.text or "").strip()
            if t.endswith("用例"):
                case_type = t.replace("用例", "")
            elif t.startswith("P"):
                priority = t

    # 子节点：按 labels 区分字段
    pre_condition = ""
    test_data = ""
    expected = ""
    steps: list[str] = []
    step_exps: list[str] = []
    children = topic.find(f"{ns}children/{ns}topics")
    if children is not None:
        for child in children.findall(f"{ns}topic"):
            clabels = child.find(f"{ns}labels")
            label_text = ""
            if clabels is not None:
                first = clabels.find(f"{ns}label")
                if first is not None:
                    label_text = (first.text or "").strip()
            ctitle_el = child.find(f"{ns}title")
            ctitle = (ctitle_el.text or "").strip() if ctitle_el is not None else ""
            if label_text == "前置条件":
                pre_condition = ctitle
            elif label_text == "测试数据":
                test_data = ctitle
            elif label_text == "操作步骤":
                steps.append(_strip_step_num(ctitle))
                # 每个操作步骤下可能有对应的预期结果子节点（V2.7 新格式）
                exp_child = child.find(f"{ns}children/{ns}topics")
                step_exp = ""
                if exp_child is not None:
                    for ec in exp_child.findall(f"{ns}topic"):
                        elabels = ec.find(f"{ns}labels")
                        if elabels is not None:
                            elabel = elabels.find(f"{ns}label")
                            if elabel is not None and (elabel.text or "").strip() == "预期结果":
                                etitle = ec.find(f"{ns}title")
                                step_exp = (etitle.text or "").strip() if etitle is not None else ""
                                break
                step_exps.append(step_exp)
            elif label_text == "预期结果":
                expected = ctitle

    # 整体预期兜底：xmind 无独立整体节点时，取最后一步预期作为 expected（保证非空）
    if not expected and step_exps:
        expected = step_exps[-1]

    return TestCase(
        title=clean_title or "未命名用例",
        module="",  # 模块由外层填充
        case_type=_to_case_type(case_type),
        priority=_to_priority(priority),
        pre_condition=pre_condition,
        steps=steps,
        step_expectations=align_step_expectations(steps, step_exps, expected),
        expected=expected,
        test_data=test_data or None,
    )


def import_xmind(path: str) -> list[TestCase]:
    """解析 xmind 文件（兼容 XMind 8 legacy XML 和新版 JSON）。"""
    try:
        zf = zipfile.ZipFile(path, "r")
    except (zipfile.BadZipFile, OSError) as e:
        raise ImportError(f"xmind 文件损坏或不是合法 zip：{e}") from e
    names = zf.namelist()
    try:
        if "content.json" in names:
            return _import_xmind_json(zf)
        if "content.xml" in names:
            return _import_xmind_xml(zf)
        raise ImportError("xmind 中未找到 content.xml 或 content.json")
    finally:
        zf.close()


def _import_xmind_xml(zf: zipfile.ZipFile) -> list[TestCase]:
    """解析 XMind 8 legacy XML 格式。"""
    with zf.open("content.xml") as f:
        tree = ET.parse(f)
    root = tree.getroot()
    # 命名空间
    ns = ""
    m = "{urn:xmind:xmap:xmlns:content:2.0}"
    if root.tag.startswith("{"):
        ns = m
    # 结构：xmap-content > sheet > topic(root) > children/topics > topic(模块) > children/topics > topic(用例)
    sheet = root.find(f"{ns}sheet")
    if sheet is None:
        raise ImportError("xmind XML 中未找到 sheet")
    root_topic = sheet.find(f"{ns}topic")
    if root_topic is None:
        raise ImportError("xmind XML 中未找到根 topic")

    cases: list[TestCase] = []
    modules = root_topic.find(f"{ns}children/{ns}topics")
    if modules is None:
        return cases
    for mod_topic in modules.findall(f"{ns}topic"):
        mod_title_el = mod_topic.find(f"{ns}title")
        module = (mod_title_el.text or "").strip() if mod_title_el is not None else "未分类"
        case_topics = mod_topic.find(f"{ns}children/{ns}topics")
        if case_topics is None:
            continue
        for case_topic in case_topics.findall(f"{ns}topic"):
            c = _parse_case_topic_xml(case_topic, ns)
            c.module = module
            cases.append(c)
    if not cases:
        raise ImportError("xmind 中未解析到有效用例（结构可能不是平台导出格式）")
    return cases


def _import_xmind_json(zf: zipfile.ZipFile) -> list[TestCase]:
    """解析新版 XMind（2020+）JSON 格式。"""
    with zf.open("content.json") as f:
        data = json.load(f)
    if not isinstance(data, list) or not data:
        raise ImportError("xmind content.json 格式异常")
    sheet = data[0]
    root_topic = sheet.get("rootTopic") or {}
    cases: list[TestCase] = []
    # 模块节点
    modules = (root_topic.get("children") or {}).get("attached") or []
    for mod in modules:
        module = mod.get("title", "未分类")
        case_topics = (mod.get("children") or {}).get("attached") or []
        for ct in case_topics:
            c = _parse_case_topic_json(ct)
            c.module = module
            cases.append(c)
    if not cases:
        raise ImportError("xmind 中未解析到有效用例（结构可能不是平台导出格式）")
    return cases


def _parse_case_topic_json(topic: dict) -> TestCase:
    """从 JSON topic 还原一个 TestCase。"""
    import re
    raw_title = topic.get("title", "")
    clean_title = re.sub(r"^\d+\.\s*", "", raw_title)
    clean_title = re.sub(r"->.*$", "", clean_title).strip()

    labels = topic.get("labels") or []
    case_type = "正向"
    priority = "P1"
    for lb in labels:
        if str(lb).endswith("用例"):
            case_type = str(lb).replace("用例", "")
        elif str(lb).startswith("P"):
            priority = str(lb)

    pre_condition = ""
    test_data = ""
    expected = ""
    steps: list[str] = []
    step_exps: list[str] = []
    children = (topic.get("children") or {}).get("attached") or []
    for child in children:
        clabels = child.get("labels") or []
        label_text = str(clabels[0]) if clabels else ""
        ctitle = child.get("title", "")
        if label_text == "前置条件":
            pre_condition = ctitle
        elif label_text == "测试数据":
            test_data = ctitle
        elif label_text == "操作步骤":
            steps.append(_strip_step_num(ctitle))
            # 每个操作步骤下可能有对应的预期结果子节点（V2.7 新格式）
            step_exp = ""
            exp_children = (child.get("children") or {}).get("attached") or []
            for ec in exp_children:
                elabels = ec.get("labels") or []
                if elabels and str(elabels[0]) == "预期结果":
                    step_exp = ec.get("title", "")
                    break
            step_exps.append(step_exp)
        elif label_text == "预期结果":
            expected = ctitle

    # 整体预期兜底：xmind 无独立整体节点时，取最后一步预期作为 expected（保证非空）
    if not expected and step_exps:
        expected = step_exps[-1]

    return TestCase(
        title=clean_title or "未命名用例",
        module="",
        case_type=_to_case_type(case_type),
        priority=_to_priority(priority),
        pre_condition=pre_condition,
        steps=steps,
        step_expectations=align_step_expectations(steps, step_exps, expected),
        expected=expected,
        test_data=test_data or None,
    )


# ============ 统一入口 ============

def import_cases(path: str, ext: str) -> tuple[list[TestCase], dict]:
    """根据扩展名分发解析器，返回 (cases, summary)。"""
    ext = ext.lower().lstrip(".")
    if ext == "json":
        cases = import_json(path)
    elif ext == "xlsx":
        cases = import_xlsx(path)
    elif ext == "xmind":
        cases = import_xmind(path)
    else:
        raise ImportError(f"不支持的用例文件格式：.{ext}（支持 xmind/xlsx/json）")
    # 重编号
    for i, c in enumerate(cases, 1):
        c.case_id = f"TC-{i:03d}"
    summary = {
        "total": len(cases),
        "modules": dict(Counter(c.module or "未分类" for c in cases)),
        "by_type": dict(Counter(_norm_type(c.case_type) for c in cases)),
        "by_priority": dict(Counter(str(c.priority) for c in cases)),
    }
    return cases, summary
